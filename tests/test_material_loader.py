"""Extraction must retain teaching structure and late visual-only pages."""
import tempfile
import importlib.util
import unittest
from pathlib import Path
from unittest.mock import patch

from core.material_loader import TeachingMaterialLoader


class LoaderTests(unittest.TestCase):
    def test_word_preserves_options_breaks_tables_emphasis_and_footer(self):
        from docx import Document
        from docx.shared import Inches
        from PIL import Image
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root/'figure.png'
            Image.new('RGB', (32,32), 'white').save(image)
            document = Document()
            document.add_heading('Learning objectives', 1)
            paragraph = document.add_paragraph('Classify health models.')
            paragraph.add_run(' Must know!').bold = True
            paragraph.add_run(' Ordinary text.').bold = False
            question = document.add_paragraph('Which model applies?')
            question.add_run().add_break()
            question.add_run('A) Clinical model')
            table = document.add_table(rows=1, cols=2)
            table.cell(0,0).text = 'Clinical model'
            table.cell(0,1).text = 'Absence of disease'
            table.cell(0,1).paragraphs[0].add_run().add_picture(str(image), width=Inches(.2))
            document.sections[0].footer.paragraphs[0].text = 'Family theories will not be tested.'
            path = root/'notes.docx'; document.save(path)
            with patch('core.material_loader.describe_visual', return_value='[Table figure]\nDiagram labels: Clinical model') as read:
                text = '\n'.join(TeachingMaterialLoader(str(path))._word())
            self.assertIn('[Heading: Learning objectives]', text)
            self.assertIn('Which model applies?\nA) Clinical model', text)
            self.assertIn('Clinical model | Absence of disease', text)
            self.assertIn('[Author emphasis:  Must know!]', text)
            self.assertNotIn('[Author emphasis:  Ordinary text.', text)
            self.assertIn('Family theories will not be tested.', text)
            self.assertIn('Diagram labels: Clinical model', text)
            read.assert_called_once()

    @unittest.skipUnless(importlib.util.find_spec('fitz'), 'PyMuPDF is installed in the app environment, not this bundled test runtime')
    def test_pdf_reads_an_image_only_page_after_text_pages(self):
        import fitz
        from PIL import Image
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder); image = root/'figure.png'
            Image.new('RGB',(80,80),'white').save(image)
            document = fitz.open()
            for index in range(4):
                page = document.new_page()
                page.insert_text((40,40), f'Learning objective {index}. ' + 'Recognize the taught health framework. '*5)
            page = document.new_page(); page.insert_image(fitz.Rect(20,20,100,100),filename=str(image))
            path = root/'mixed.pdf'; document.save(path); document.close()
            with patch('core.material_loader.describe_visual',return_value='Final page: MUST KNOW family assessment') as read:
                pages = TeachingMaterialLoader(str(path))._pdf()
            self.assertEqual(len(pages),5)
            self.assertIn('Final page: MUST KNOW family assessment',pages[-1])
            self.assertEqual(read.call_args.args[-1], 'PDF page 5')

    def test_slides_keep_speaker_notes_and_tables(self):
        from pptx import Presentation
        from pptx.util import Inches
        with tempfile.TemporaryDirectory() as folder:
            deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = 'Learning objective: use SBAR'
            table = slide.shapes.add_table(1,2,Inches(1),Inches(2), Inches(5), Inches(1)).table
            table.cell(0,0).text = 'S'; table.cell(0,1).text = 'Situation'
            slide.notes_slide.notes_text_frame.text = 'Exam instruction: recognize each SBAR component.'
            path = Path(folder)/'notes.pptx'; deck.save(path)
            text = '\n'.join(TeachingMaterialLoader(str(path))._slides())
            self.assertIn('Learning objective: use SBAR',text)
            self.assertIn('S | Situation',text)
            self.assertIn('Exam instruction: recognize each SBAR component.',text)

    def test_plain_text_pages_with_table_rules_are_not_sent_as_images(self):
        # 2026-10-08: a 4-page handout sent every page to the vision model
        # because table borders count as drawings; 34s before analysis began.
        import fitz
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder)/'handout.pdf'
            document = fitz.open()
            for index in range(3):
                page = document.new_page()
                page.insert_textbox(fitz.Rect(40, 40, 560, 780), ' '.join(['Airway breathing circulation disability exposure.'] * 20), fontsize=9)
                page.draw_rect(fitz.Rect(30, 30, 560, 400))          # a table border: a drawing, not an image
            sparse = document.new_page()
            sparse.draw_line((40, 40), (300, 300))                     # a diagram with almost no text
            sparse.insert_text((40, 320), 'Fig 1')
            document.save(path); document.close()
            with patch('core.material_loader.describe_visual', return_value='diagram described') as read:
                pages = TeachingMaterialLoader(str(path))._pdf()
            self.assertEqual(read.call_count, 1)
            self.assertEqual(read.call_args.args[2], 'PDF page 4')
            self.assertIn('diagram described', pages[3])
            self.assertTrue(all('\x00' not in page for page in pages))

    def test_figures_are_described_concurrently_and_land_in_order(self):
        import threading, time
        from core import material_loader
        active, peak, lock = [0], [0], threading.Lock()
        def slow(data, mime, label, **kwargs):
            with lock:
                active[0] += 1; peak[0] = max(peak[0], active[0])
            time.sleep(0.05)
            with lock:
                active[0] -= 1
            return f'described {label}'
        loader = TeachingMaterialLoader('deck.pptx')
        parts = [f'[Slide {i}] ' + loader._queue_visual(b'x', 'image/png', f'figure {i}') for i in range(6)]
        with patch('core.material_loader.describe_visual', side_effect=slow):
            resolved = loader._resolve_visuals(parts)
        self.assertEqual(resolved, [f'[Slide {i}] described figure {i}' for i in range(6)])
        self.assertGreater(peak[0], 1)
        self.assertLessEqual(peak[0], material_loader.VISUAL_WORKERS)

    def test_a_header_logo_does_not_send_a_text_page_to_the_vision_model(self):
        import fitz
        from PIL import Image
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder); logo = root/'logo.png'
            Image.new('RGB', (252, 126), 'green').save(logo)
            body = ' '.join(['Primary survey: catastrophic haemorrhage first, then airway.'] * 20)
            document = fitz.open()
            header = document.new_page()
            header.insert_image(fitz.Rect(40, 30, 160, 90), filename=str(logo))       # letterhead logo
            header.insert_textbox(fitz.Rect(40, 120, 560, 780), body, fontsize=9)
            figure = document.new_page()
            figure.insert_textbox(fitz.Rect(40, 40, 560, 300), body, fontsize=9)
            figure.insert_image(fitz.Rect(200, 360, 320, 420), filename=str(logo))    # same size, mid-page
            path = root/'handout.pdf'; document.save(path); document.close()
            with patch('core.material_loader.describe_visual', return_value='figure described') as read:
                pages = TeachingMaterialLoader(str(path))._pdf()
            self.assertEqual(read.call_count, 1)
            self.assertEqual(read.call_args.args[2], 'PDF page 2')
            self.assertNotIn('figure described', pages[0])
            self.assertIn('figure described', pages[1])


if __name__ == '__main__': unittest.main()
